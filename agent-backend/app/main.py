import json
import os
import re
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parents[2]
STATIC_DIR = Path(__file__).resolve().parent / "static"
CATALOG_XLSX_PATH = Path(os.getenv("CATALOG_XLSX_PATH", str(ROOT / "AI穿搭导购_商品更新源.xlsx")))
DATABASE_PATH = Path(os.getenv("DATABASE_PATH", str(ROOT / "agent-backend" / "data" / "agent.db")))
DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
ALLOWED_ORIGINS = [item.strip() for item in os.getenv("ALLOWED_ORIGINS", "http://localhost:3000").split(",") if item.strip()]
PROFILE_ID_PATTERN = re.compile(r"^[A-Za-z0-9_-]{8,80}$")

app = FastAPI(title="衣橱顾问 Agent API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=False,
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["Content-Type", "X-Request-ID"],
)


class ChatRequest(BaseModel):
    message: str = Field(min_length=1, max_length=500)
    session_id: str | None = Field(default=None, max_length=100)
    profile_id: str | None = Field(default=None, max_length=80)


class ImportRequest(BaseModel):
    source_path: str | None = None


class ProfileInput(BaseModel):
    common_scenarios: list[str] = Field(default_factory=list, max_length=10)
    budget_range: str | None = Field(default=None, max_length=40)
    usual_size: str | None = Field(default=None, max_length=20)
    preferred_styles: list[str] = Field(default_factory=list, max_length=10)
    preferred_colors: list[str] = Field(default_factory=list, max_length=10)
    avoided_colors: list[str] = Field(default_factory=list, max_length=10)
    comfort_preferences: list[str] = Field(default_factory=list, max_length=10)
    wardrobe_goal: str | None = Field(default=None, max_length=80)


class WardrobeItemInput(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    category: str = Field(min_length=1, max_length=30)
    color: str | None = Field(default=None, max_length=30)
    styles: list[str] = Field(default_factory=list, max_length=8)
    scenarios: list[str] = Field(default_factory=list, max_length=8)
    fit: str | None = Field(default=None, max_length=40)
    size: str | None = Field(default=None, max_length=20)
    ownership_status: str = Field(default="已拥有", max_length=20)
    image_url: str | None = Field(default=None, max_length=500)
    idle: bool = False


class CalendarEntryInput(BaseModel):
    wardrobe_item_ids: list[str] = Field(default_factory=list, max_length=10)
    outfit_title: str | None = Field(default=None, max_length=100)
    scene: str | None = Field(default=None, max_length=30)
    weather_feeling: str | None = Field(default=None, max_length=120)
    comfort_rating: int | None = Field(default=None, ge=1, le=5)
    satisfaction_rating: int | None = Field(default=None, ge=1, le=5)
    note: str | None = Field(default=None, max_length=500)


@contextmanager
def db_connection():
    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row
    try:
        yield connection
        connection.commit()
    finally:
        connection.close()


def now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def validate_profile_id(profile_id: str) -> str:
    if not PROFILE_ID_PATTERN.fullmatch(profile_id):
        raise HTTPException(status_code=422, detail="profile_id 格式无效")
    return profile_id


def json_value(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False)


def json_load(value: str | None, fallback: Any) -> Any:
    if not value:
        return fallback
    try:
        return json.loads(value)
    except json.JSONDecodeError:
        return fallback


def init_db() -> None:
    with db_connection() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS catalog_items (
                product_id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                category TEXT NOT NULL,
                subcategory TEXT,
                styles TEXT,
                scenarios TEXT,
                color TEXT,
                fit TEXT,
                slim_tag TEXT,
                size_range TEXT,
                height_range TEXT,
                price REAL NOT NULL,
                inventory INTEGER NOT NULL,
                sale_status TEXT NOT NULL,
                recommendable TEXT NOT NULL,
                image_url TEXT,
                source TEXT,
                updated_at TEXT,
                note TEXT
            );
            CREATE TABLE IF NOT EXISTS sessions (
                session_id TEXT PRIMARY KEY,
                state_json TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS traces (
                trace_id TEXT PRIMARY KEY,
                session_id TEXT NOT NULL,
                user_message TEXT NOT NULL,
                state_json TEXT NOT NULL,
                tool_summary TEXT NOT NULL,
                response_json TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS user_profiles (
                profile_id TEXT PRIMARY KEY,
                profile_json TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS wardrobe_items (
                item_id TEXT PRIMARY KEY,
                profile_id TEXT NOT NULL,
                name TEXT NOT NULL,
                category TEXT NOT NULL,
                color TEXT,
                styles_json TEXT NOT NULL,
                scenarios_json TEXT NOT NULL,
                fit TEXT,
                size TEXT,
                ownership_status TEXT NOT NULL,
                image_url TEXT,
                idle INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_wardrobe_profile ON wardrobe_items(profile_id);
            CREATE TABLE IF NOT EXISTS outfit_calendar_entries (
                entry_id TEXT PRIMARY KEY,
                profile_id TEXT NOT NULL,
                wear_date TEXT NOT NULL,
                wardrobe_item_ids_json TEXT NOT NULL,
                outfit_title TEXT,
                scene TEXT,
                weather_feeling TEXT,
                comfort_rating INTEGER,
                satisfaction_rating INTEGER,
                note TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(profile_id, wear_date)
            );
            CREATE INDEX IF NOT EXISTS idx_calendar_profile_month ON outfit_calendar_entries(profile_id, wear_date);
            """
        )


def value(row: tuple[Any, ...], index: int, default: Any = "") -> Any:
    item = row[index] if len(row) > index else default
    return default if item is None else item


def import_catalog(source_path: Path) -> dict[str, Any]:
    if not source_path.exists() or source_path.suffix.lower() != ".xlsx":
        raise HTTPException(status_code=422, detail="商品更新源必须是可访问的 .xlsx 文件")
    workbook = load_workbook(source_path, data_only=False, read_only=True)
    if "商品主数据" not in workbook.sheetnames:
        raise HTTPException(status_code=422, detail="工作簿缺少“商品主数据”工作表")
    rows = list(workbook["商品主数据"].iter_rows(min_row=2, values_only=True))
    products: list[tuple[Any, ...]] = []
    for row in rows:
        product_id = str(value(row, 1)).strip()
        if not product_id:
            continue
        try:
            price = float(value(row, 12, 0))
            inventory = int(value(row, 13, 0))
        except (TypeError, ValueError) as error:
            raise HTTPException(status_code=422, detail=f"商品 {product_id} 的价格或库存无效") from error
        sale_status = str(value(row, 14))
        products.append((
            product_id, str(value(row, 2)), str(value(row, 3)), str(value(row, 4)),
            str(value(row, 5)), str(value(row, 6)), str(value(row, 7)), str(value(row, 8)),
            str(value(row, 9)), str(value(row, 10)), str(value(row, 11)), price, inventory,
            sale_status, "是" if inventory > 0 and sale_status == "在售" else "否", str(value(row, 16)),
            str(value(row, 17)), str(value(row, 18)), str(value(row, 19)),
        ))
    if len(products) != 70:
        raise HTTPException(status_code=422, detail=f"商品主数据应为 70 条，当前读取到 {len(products)} 条")
    if len({item[0] for item in products}) != len(products):
        raise HTTPException(status_code=422, detail="商品 ID 存在重复，拒绝导入")
    with db_connection() as db:
        db.execute("DELETE FROM catalog_items")
        db.executemany(
            """INSERT INTO catalog_items (
                product_id,name,category,subcategory,styles,scenarios,color,fit,slim_tag,
                size_range,height_range,price,inventory,sale_status,recommendable,image_url,
                source,updated_at,note
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            products,
        )
    return {"source": str(source_path), "imported": len(products), "updated_at": now()}


def default_state() -> dict[str, Any]:
    return {"scenario": None, "budget": None, "style": None, "height": None, "size": None,
            "slim": False, "excluded_categories": [], "question_count": 0}


def get_session(session_id: str) -> dict[str, Any]:
    with db_connection() as db:
        record = db.execute("SELECT state_json FROM sessions WHERE session_id = ?", (session_id,)).fetchone()
    return json_load(record["state_json"], default_state()) if record else default_state()


def save_session(session_id: str, state: dict[str, Any]) -> None:
    with db_connection() as db:
        db.execute(
            "INSERT INTO sessions (session_id,state_json,updated_at) VALUES (?,?,?) "
            "ON CONFLICT(session_id) DO UPDATE SET state_json=excluded.state_json,updated_at=excluded.updated_at",
            (session_id, json_value(state), now()),
        )


def parse_conditions(message: str, state: dict[str, Any]) -> tuple[dict[str, Any], list[str], list[str]]:
    changes: list[str] = []
    normalized = message.replace("￥", "¥").replace("以内", "内")
    matched_scenarios = [item for item in ["通勤", "约会", "旅行", "面试", "日常", "聚会"] if item in normalized]
    if len(matched_scenarios) == 1 and state.get("scenario") != matched_scenarios[0]:
        state["scenario"] = matched_scenarios[0]
        changes.append(f"场景更新为{matched_scenarios[0]}")
    for style in ["简约", "法式", "运动", "休闲", "复古"]:
        if style in normalized and state.get("style") != style:
            state["style"] = style
            changes.append(f"风格更新为{style}")
    budget_match = re.search(r"(?:¥|预算|不超过|控制在)?\s*(\d{2,4})\s*(?:元|块|¥|以内|内)", normalized)
    if budget_match:
        state["budget"] = int(budget_match.group(1))
        changes.append(f"预算更新为¥{state['budget']}")
    elif any(token in normalized for token in ["再便宜", "便宜一点", "更便宜"]):
        state["budget"] = max(200, (state.get("budget") or 500) - 100)
        changes.append(f"预算下调至¥{state['budget']}")
    height_match = re.search(r"(1[4-8]\d)\s*(?:cm|厘米|公分)?", normalized, flags=re.I)
    if height_match:
        state["height"] = int(height_match.group(1))
        changes.append(f"身高更新为{state['height']}cm")
    size_match = re.search(r"(?<![A-Z])(S|M|L|XL)(?:码|号)?", normalized, flags=re.I)
    if size_match:
        state["size"] = size_match.group(1).upper()
        changes.append(f"尺码更新为{state['size']}码")
    if "更显瘦" in normalized or "显瘦" in normalized:
        state["slim"] = True
        changes.append("已优先筛选显瘦版型")
    if "不要裙子" in normalized or "不穿裙子" in normalized:
        state["excluded_categories"] = list({*state["excluded_categories"], "半身裙", "连衣裙"})
        changes.append("已排除半身裙和连衣裙")
    return state, changes, matched_scenarios


def missing_fields(state: dict[str, Any]) -> list[str]:
    return [field for field in ["scenario", "budget", "style"] if not state.get(field)]


def ask_question(state: dict[str, Any]) -> str:
    missing = missing_fields(state)
    if "scenario" in missing:
        return "这套穿搭主要用于什么场景？例如通勤、约会、旅行或面试。"
    if "budget" in missing:
        return "你的总预算大约是多少？例如 ¥300 内或 ¥500 内。"
    return "你更偏好简约、法式、运动、休闲还是复古？"


def query_products(state: dict[str, Any]) -> list[dict[str, Any]]:
    clauses = ["sale_status = '在售'", "recommendable = '是'", "inventory > 0"]
    params: list[Any] = []
    if state.get("size"):
        clauses.append("size_range LIKE ?")
        params.append(f"%{state['size']}%")
    for category in state.get("excluded_categories", []):
        clauses.append("category != ?")
        params.append(category)
    if state.get("scenario"):
        clauses.append("scenarios LIKE ?")
        params.append(f"%{state['scenario']}%")
    if state.get("style"):
        clauses.append("styles LIKE ?")
        params.append(f"%{state['style']}%")
    where = " AND ".join(clauses)
    order = "CASE WHEN slim_tag = '显瘦' THEN 0 ELSE 1 END, price ASC" if state.get("slim") else "price ASC"
    with db_connection() as db:
        records = db.execute(f"SELECT * FROM catalog_items WHERE {where} ORDER BY {order}", params).fetchall()
    return [dict(record) for record in records]


def build_outfits(items: list[dict[str, Any]], state: dict[str, Any]) -> list[dict[str, Any]]:
    budget = state.get("budget") or 9999
    patterns = [["上装", "裤装", "鞋履", "配饰"], ["连衣裙", "鞋履", "配饰"], ["上装", "半身裙", "鞋履", "配饰"]]
    outfits: list[dict[str, Any]] = []
    used: set[str] = set()
    for pattern in patterns:
        if any(category in state.get("excluded_categories", []) for category in pattern):
            continue
        picks: list[dict[str, Any]] = []
        local_used = set(used)
        for category in pattern:
            candidate = next((item for item in items if item["category"] == category and item["product_id"] not in local_used), None)
            if not candidate:
                picks = []
                break
            picks.append(candidate)
            local_used.add(candidate["product_id"])
        total = sum(item["price"] for item in picks)
        if picks and total <= budget:
            used.update(local_used)
            score = 85 + (5 if state.get("slim") else 0) + (3 if state.get("height") and state.get("size") else 0)
            score = min(98, score)
            outfits.append({
                "name": f"{state.get('scenario') or '日常'} {state.get('style') or '轻松'}搭配 {len(outfits) + 1}",
                "items": [{key: product[key] for key in ["product_id", "name", "category", "price", "image_url", "fit", "slim_tag", "size_range"]} for product in picks],
                "total_price": total,
                "reason": "已严格校验在售、库存、预算、尺码与禁忌品类。" + (" 同时优先选择显瘦版型。" if state.get("slim") else ""),
                "size_advice": f"建议优先试穿 {state.get('size') or '常穿'} 码；具体以商品尺码表为准。",
                "match_score": score,
                "match_level": "高度适配" if score >= 85 else "较适配",
            })
        if len(outfits) == 3:
            break
    return outfits


def state_summary(state: dict[str, Any]) -> dict[str, Any]:
    return {"场景": state.get("scenario"), "预算": f"¥{state['budget']}" if state.get("budget") else None,
            "风格": state.get("style"), "身高": f"{state['height']}cm" if state.get("height") else None,
            "尺码": state.get("size"), "显瘦优先": "是" if state.get("slim") else "否",
            "排除品类": state.get("excluded_categories"), "已追问轮次": state.get("question_count")}


def profile_payload(profile_id: str) -> dict[str, Any]:
    with db_connection() as db:
        row = db.execute("SELECT profile_json,updated_at FROM user_profiles WHERE profile_id = ?", (profile_id,)).fetchone()
    return {"profile_id": profile_id, "profile": json_load(row["profile_json"], {}) if row else {}, "updated_at": row["updated_at"] if row else None}


def wardrobe_payload(row: sqlite3.Row) -> dict[str, Any]:
    item = dict(row)
    item["styles"] = json_load(item.pop("styles_json"), [])
    item["scenarios"] = json_load(item.pop("scenarios_json"), [])
    item["idle"] = bool(item["idle"])
    return item


@app.on_event("startup")
def startup() -> None:
    init_db()
    with db_connection() as db:
        count = db.execute("SELECT COUNT(*) AS count FROM catalog_items").fetchone()["count"]
    if count == 0 and CATALOG_XLSX_PATH.exists():
        import_catalog(CATALOG_XLSX_PATH)


@app.get("/")
def home() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/assets/{filename}")
def assets(filename: str) -> FileResponse:
    target = STATIC_DIR / filename
    if not target.exists() or target.parent != STATIC_DIR:
        raise HTTPException(status_code=404, detail="静态资源未找到")
    return FileResponse(target)


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/ready")
def ready() -> dict[str, Any]:
    with db_connection() as db:
        count = db.execute("SELECT COUNT(*) AS count FROM catalog_items").fetchone()["count"]
    return {"status": "ok" if count == 70 else "degraded", "catalog_items": count}


@app.post("/api/catalog/import-xlsx")
def refresh_catalog(request: ImportRequest) -> dict[str, Any]:
    source = Path(request.source_path) if request.source_path else CATALOG_XLSX_PATH
    return import_catalog(source)


@app.post("/api/chat")
def chat(request: ChatRequest) -> dict[str, Any]:
    session_id = request.session_id or str(uuid.uuid4())
    trace_id = str(uuid.uuid4())
    state = get_session(session_id)
    state, changes, candidate_scenes = parse_conditions(request.message, state)
    if len(candidate_scenes) >= 2:
        payload = {"trace_id": trace_id, "session_id": session_id, "type": "scene_conflict",
                   "message": f"你同时提到了{'和'.join(candidate_scenes)}，这次更偏向哪个场景？",
                   "candidate_scenes": candidate_scenes, "changes": [], "state": state_summary(state), "outfits": []}
    elif missing_fields(state) and state["question_count"] < 3:
        state["question_count"] += 1
        payload = {"trace_id": trace_id, "session_id": session_id, "type": "follow_up", "message": ask_question(state),
                   "changes": changes, "state": state_summary(state), "outfits": []}
    else:
        outfits = build_outfits(query_products(state), state)
        if outfits:
            payload = {"trace_id": trace_id, "session_id": session_id, "type": "recommendation",
                       "message": f"已为你找到 {len(outfits)} 套完整穿搭。", "changes": changes,
                       "state": state_summary(state), "outfits": outfits}
        else:
            payload = {"trace_id": trace_id, "session_id": session_id, "type": "no_match",
                       "message": "当前条件下暂无完整穿搭。可尝试增加预算或放宽风格。", "changes": changes,
                       "state": state_summary(state), "outfits": [],
                       "blocking_reasons": ["当前可售商品无法同时组成符合预算与尺码的完整搭配"],
                       "actions": ["增加 ¥100 预算", "放宽风格", "换一个场景"]}
    save_session(session_id, state)
    with db_connection() as db:
        db.execute("INSERT INTO traces (trace_id,session_id,user_message,state_json,tool_summary,response_json,created_at) VALUES (?,?,?,?,?,?,?)",
                   (trace_id, session_id, request.message, json_value(state), payload["type"], json_value(payload), now()))
    return payload


@app.get("/api/traces/{trace_id}")
def get_trace(trace_id: str) -> dict[str, Any]:
    with db_connection() as db:
        record = db.execute("SELECT * FROM traces WHERE trace_id = ?", (trace_id,)).fetchone()
    if not record:
        raise HTTPException(status_code=404, detail="trace 未找到")
    return dict(record)


@app.get("/api/profiles/{profile_id}")
def get_profile(profile_id: str) -> dict[str, Any]:
    return profile_payload(validate_profile_id(profile_id))


@app.put("/api/profiles/{profile_id}")
def save_profile(profile_id: str, profile: ProfileInput) -> dict[str, Any]:
    profile_id = validate_profile_id(profile_id)
    with db_connection() as db:
        db.execute("INSERT INTO user_profiles (profile_id,profile_json,created_at,updated_at) VALUES (?,?,?,?) "
                   "ON CONFLICT(profile_id) DO UPDATE SET profile_json=excluded.profile_json,updated_at=excluded.updated_at",
                   (profile_id, json_value(profile.model_dump()), now(), now()))
    return profile_payload(profile_id)


@app.delete("/api/profiles/{profile_id}")
def delete_profile(profile_id: str) -> dict[str, str]:
    profile_id = validate_profile_id(profile_id)
    with db_connection() as db:
        db.execute("DELETE FROM user_profiles WHERE profile_id = ?", (profile_id,))
    return {"status": "deleted"}


@app.get("/api/profiles/{profile_id}/wardrobe")
def list_wardrobe(profile_id: str) -> dict[str, Any]:
    profile_id = validate_profile_id(profile_id)
    with db_connection() as db:
        rows = db.execute("SELECT item.*, COUNT(calendar.entry_id) AS wear_count FROM wardrobe_items item "
                          "LEFT JOIN outfit_calendar_entries calendar ON calendar.wardrobe_item_ids_json LIKE '%' || item.item_id || '%' "
                          "WHERE item.profile_id = ? GROUP BY item.item_id ORDER BY item.updated_at DESC", (profile_id,)).fetchall()
    return {"items": [{**wardrobe_payload(row), "wear_count": row["wear_count"]} for row in rows]}


@app.post("/api/profiles/{profile_id}/wardrobe")
def create_wardrobe_item(profile_id: str, item: WardrobeItemInput) -> dict[str, Any]:
    profile_id = validate_profile_id(profile_id)
    item_id = str(uuid.uuid4())
    with db_connection() as db:
        db.execute("INSERT INTO wardrobe_items (item_id,profile_id,name,category,color,styles_json,scenarios_json,fit,size,ownership_status,image_url,idle,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                   (item_id, profile_id, item.name, item.category, item.color, json_value(item.styles), json_value(item.scenarios), item.fit, item.size, item.ownership_status, item.image_url, int(item.idle), now(), now()))
        row = db.execute("SELECT * FROM wardrobe_items WHERE item_id = ?", (item_id,)).fetchone()
    return wardrobe_payload(row)


@app.put("/api/profiles/{profile_id}/wardrobe/{item_id}")
def update_wardrobe_item(profile_id: str, item_id: str, item: WardrobeItemInput) -> dict[str, Any]:
    profile_id = validate_profile_id(profile_id)
    with db_connection() as db:
        current = db.execute("SELECT item_id FROM wardrobe_items WHERE item_id = ? AND profile_id = ?", (item_id, profile_id)).fetchone()
        if not current:
            raise HTTPException(status_code=404, detail="衣橱单品未找到")
        db.execute("UPDATE wardrobe_items SET name=?,category=?,color=?,styles_json=?,scenarios_json=?,fit=?,size=?,ownership_status=?,image_url=?,idle=?,updated_at=? WHERE item_id=?",
                   (item.name, item.category, item.color, json_value(item.styles), json_value(item.scenarios), item.fit, item.size, item.ownership_status, item.image_url, int(item.idle), now(), item_id))
        row = db.execute("SELECT * FROM wardrobe_items WHERE item_id = ?", (item_id,)).fetchone()
    return wardrobe_payload(row)


@app.delete("/api/profiles/{profile_id}/wardrobe/{item_id}")
def delete_wardrobe_item(profile_id: str, item_id: str) -> dict[str, str]:
    profile_id = validate_profile_id(profile_id)
    with db_connection() as db:
        db.execute("DELETE FROM wardrobe_items WHERE item_id = ? AND profile_id = ?", (item_id, profile_id))
    return {"status": "deleted"}


@app.get("/api/profiles/{profile_id}/calendar")
def list_calendar(profile_id: str, month: str = Query(pattern=r"^\d{4}-\d{2}$")) -> dict[str, Any]:
    profile_id = validate_profile_id(profile_id)
    with db_connection() as db:
        rows = db.execute("SELECT * FROM outfit_calendar_entries WHERE profile_id = ? AND substr(wear_date, 1, 7) = ? ORDER BY wear_date", (profile_id, month)).fetchall()
    entries = []
    for row in rows:
        entry = dict(row)
        entry["wardrobe_item_ids"] = json_load(entry.pop("wardrobe_item_ids_json"), [])
        entries.append(entry)
    return {"entries": entries}


@app.put("/api/profiles/{profile_id}/calendar/{wear_date}")
def save_calendar_entry(profile_id: str, wear_date: str, entry: CalendarEntryInput) -> dict[str, Any]:
    profile_id = validate_profile_id(profile_id)
    try:
        date.fromisoformat(wear_date)
    except ValueError as error:
        raise HTTPException(status_code=422, detail="wear_date 必须为 YYYY-MM-DD") from error
    entry_id = str(uuid.uuid4())
    with db_connection() as db:
        db.execute("INSERT INTO outfit_calendar_entries (entry_id,profile_id,wear_date,wardrobe_item_ids_json,outfit_title,scene,weather_feeling,comfort_rating,satisfaction_rating,note,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?) "
                   "ON CONFLICT(profile_id,wear_date) DO UPDATE SET wardrobe_item_ids_json=excluded.wardrobe_item_ids_json,outfit_title=excluded.outfit_title,scene=excluded.scene,weather_feeling=excluded.weather_feeling,comfort_rating=excluded.comfort_rating,satisfaction_rating=excluded.satisfaction_rating,note=excluded.note,updated_at=excluded.updated_at",
                   (entry_id, profile_id, wear_date, json_value(entry.wardrobe_item_ids), entry.outfit_title, entry.scene, entry.weather_feeling, entry.comfort_rating, entry.satisfaction_rating, entry.note, now(), now()))
        row = db.execute("SELECT * FROM outfit_calendar_entries WHERE profile_id = ? AND wear_date = ?", (profile_id, wear_date)).fetchone()
    result = dict(row)
    result["wardrobe_item_ids"] = json_load(result.pop("wardrobe_item_ids_json"), [])
    return result


@app.delete("/api/profiles/{profile_id}/calendar/{wear_date}")
def delete_calendar_entry(profile_id: str, wear_date: str) -> dict[str, str]:
    profile_id = validate_profile_id(profile_id)
    with db_connection() as db:
        db.execute("DELETE FROM outfit_calendar_entries WHERE profile_id = ? AND wear_date = ?", (profile_id, wear_date))
    return {"status": "deleted"}


@app.get("/api/profiles/{profile_id}/calendar-summary")
def calendar_summary(profile_id: str, month: str = Query(pattern=r"^\d{4}-\d{2}$")) -> dict[str, Any]:
    profile_id = validate_profile_id(profile_id)
    with db_connection() as db:
        rows = db.execute("SELECT * FROM outfit_calendar_entries WHERE profile_id = ? AND substr(wear_date,1,7) = ?", (profile_id, month)).fetchall()
        wardrobe = {row["item_id"]: row["name"] for row in db.execute("SELECT item_id,name FROM wardrobe_items WHERE profile_id = ?", (profile_id,)).fetchall()}
    counts: dict[str, int] = {}
    scenarios: set[str] = set()
    signatures: list[str] = []
    for row in rows:
        item_ids = sorted(json_load(row["wardrobe_item_ids_json"], []))
        signatures.append("|".join(item_ids))
        if row["scene"]:
            scenarios.add(row["scene"])
        for item_id in item_ids:
            counts[item_id] = counts.get(item_id, 0) + 1
    top_item_id = max(counts, key=counts.get) if counts else None
    return {"recorded_days": len(rows), "top_item": wardrobe.get(top_item_id) if top_item_id else None,
            "top_item_wears": counts.get(top_item_id, 0) if top_item_id else 0, "scene_coverage": len(scenarios),
            "repeat_ratio": round((len(signatures) - len(set(signatures))) / len(signatures), 2) if signatures else 0}
